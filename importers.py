# -*- coding: utf-8 -*-
"""
예시 식단표 가져오기 유틸
==========================
기존 구글시트 식단표를 로컬 파일로 내려받아(.xlsx/.csv/.txt) 예시 텍스트로 변환한다.
- 엑셀: 모든 시트의 비어있지 않은 셀을 행 단위로 이어붙여 텍스트화
- CSV : 행을 쉼표로 이어붙임
- TXT : 원문 그대로

레이아웃이 제각각이어도 '스타일 참고용 예시'로 쓰기 좋게 평문으로 뽑는다.
"""

from __future__ import annotations

import csv
import os

MAX_CHARS = 8000  # 예시 1개당 과도한 길이 방지(토큰 절약)


def read_table_file(path: str) -> str:
    """엑셀/CSV/텍스트 파일을 예시용 평문 텍스트로 변환."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        text = _read_xlsx(path)
    elif ext in (".csv", ".tsv"):
        text = _read_csv(path, delimiter="\t" if ext == ".tsv" else ",")
    else:  # .txt 등
        with open(path, encoding="utf-8", errors="replace") as f:
            text = f.read()
    text = text.strip()
    return text[:MAX_CHARS]


def _read_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(", ".join(cells))
        if rows:
            title = ws.title
            blocks.append(f"[{title}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(blocks)


def _read_csv(path: str, delimiter: str = ",") -> str:
    rows: list[str] = []
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        for row in csv.reader(f, delimiter=delimiter):
            cells = [c.strip() for c in row if c and c.strip()]
            if cells:
                rows.append(", ".join(cells))
    return "\n".join(rows)
