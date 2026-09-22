# -*- coding: utf-8 -*-
"""
과거 식단 기록(이력)
=====================
‘몇 월 며칠에 실제로 무엇을 냈는지’를 그대로 쌓아 두는 곳.
- 다음 달 식단을 만들 때 프롬프트에 참고 자료로 싣고,
- 달 경계(지난달 말 ~ 이번달 초) 간격 검증에도 쓰고,
- 기록이 충분히 쌓이면 ‘과거에 냈던 메뉴만’ 으로 만들 수도 있다(MODE_ONLY).

저장: config_dir()/history.json  (store.DATA_FILES 에 포함 → GitHub 동기화 대상)
형식: {"2026-07": {"1": {"중식": [6개], "석식": [6개]}, ...}, ...}
      · JSON 키라서 연-월·일 모두 문자열. 항목은 항상 6칸으로 맞춘다.
      · 여기 넣는 건 '실제로 낸 식단'만 — 만들다 만 초안은 넣지 않는다.

⚠ store 를 import 하지 않는다(store 가 이 모듈을 import 하므로 순환이 된다).
   풀에 반영하는 일은 store.add_missing_from_history() 쪽에 둔다.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date

import app_config
import layouts
from layouts import ITEMS_PER_MEAL, SECTIONS, MonthData

HISTORY_FILE = "history.json"

# 과거 기록 활용 방식(settings.json 의 history_mode)
MODE_OFF, MODE_REF, MODE_ONLY = "off", "ref", "only"
MODES = (MODE_OFF, MODE_REF, MODE_ONLY)
MODE_LABELS = {
    MODE_OFF: "참고 안 함 (지금까지처럼)",
    MODE_REF: "과거 식단을 참고해서 만들기",
    MODE_ONLY: "과거에 냈던 메뉴만으로 만들기",
}

# 한 끼 6칸에서 각 자리의 역할. 4번(김치류)은 이름을 모으지 않는다.
IDX_MAIN = 0
IDX_SIDES = (1, 2, 3)
IDX_SOUP = 5

# 프롬프트에 싣는 최근 실제 식단 개월 수 / 메인 사용현황 줄 수
PROMPT_MONTHS = 2
PROMPT_STAT_ROWS = 60

# 달 경계 검증에 쓸 '지난달 꼬리' 길이(일)
TAIL_DAYS = 10

_MKEY_RE = re.compile(r"^(\d{4})-(\d{1,2})$")


# ----------------------------------------------------------------------------
# 저장/로드
# ----------------------------------------------------------------------------

def _path():
    return app_config.config_dir() / HISTORY_FILE


def load() -> dict:
    p = _path()
    if not p.exists():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    return data if isinstance(data, dict) else {}


def save(h: dict) -> None:
    d = app_config.config_dir()
    d.mkdir(parents=True, exist_ok=True)
    _path().write_text(json.dumps(h, ensure_ascii=False, indent=2), encoding="utf-8")


def mkey(year: int, month: int) -> str:
    return f"{year:04d}-{month:02d}"


def parse_mkey(key: str) -> tuple[int, int] | None:
    m = _MKEY_RE.match(str(key).strip())
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    return (y, mo) if 1 <= mo <= 12 else None


def _pad(items) -> list[str]:
    """어떤 길이로 들어와도 6칸으로 맞춘다(표/엑셀/모델 출력이 제각각이므로)."""
    out = [str(x).strip() for x in (items or [])][:ITEMS_PER_MEAL]
    return out + [""] * (ITEMS_PER_MEAL - len(out))


# ----------------------------------------------------------------------------
# 월 단위 읽기/쓰기 (GUI 표와 같은 MonthData 를 주고받는다)
# ----------------------------------------------------------------------------

def get_month(year: int, month: int, h: dict | None = None) -> MonthData:
    h = load() if h is None else h
    md = MonthData(year=year, month=month)
    for ds, secs in (h.get(mkey(year, month)) or {}).items():
        try:
            day = int(ds)
        except (TypeError, ValueError):
            continue
        if not isinstance(secs, dict):
            continue
        for sec in SECTIONS:
            items = secs.get(sec)
            if items:
                md.meals.setdefault(day, {})[sec] = _pad(items)
    return md


def save_month(md: MonthData, merge: bool = False) -> int:
    """한 달치를 저장하고 기록된 날짜 수를 반환.

    merge=False(기본): 그 달을 md 내용으로 '교체' — 화면에서 비운 칸이 지워져야 하므로
    과거 식단 창의 저장은 이 쪽이다.
    merge=True: 기존 기록 위에 덮어쓰기만 — 엑셀 가져오기/현재 표 저장처럼
    한 달의 일부만 들어올 때 나머지 날짜를 지우지 않는다.
    """
    h = load()
    k = mkey(md.year, md.month)
    month: dict = dict(h.get(k) or {}) if merge else {}
    for day, secs in md.meals.items():
        entry = dict(month.get(str(day)) or {})
        for sec in SECTIONS:
            items = _pad(secs.get(sec) or [])
            if any(items):
                entry[sec] = items
            elif not merge:
                entry.pop(sec, None)
        if entry:
            month[str(day)] = entry
        else:
            month.pop(str(day), None)
    if month:
        h[k] = {d: month[d] for d in sorted(month, key=int)}
    else:
        h.pop(k, None)
    save(h)
    return len(month)


def delete_month(year: int, month: int) -> bool:
    h = load()
    if h.pop(mkey(year, month), None) is None:
        return False
    save(h)
    return True


def recorded_months(h: dict | None = None) -> list[tuple[int, int]]:
    h = load() if h is None else h
    out = [p for k in h if (p := parse_mkey(k)) and h.get(k)]
    return sorted(out)


def day_count(h: dict | None = None) -> int:
    h = load() if h is None else h
    return sum(len(v) for k, v in h.items() if parse_mkey(k) and isinstance(v, dict))


def summary(h: dict | None = None) -> str:
    """오른쪽 패널에 띄울 한 줄 요약."""
    h = load() if h is None else h
    months = recorded_months(h)
    if not months:
        return "아직 기록이 없어요."
    y1, m1 = months[0]
    y2, m2 = months[-1]
    span = f"{y1}.{m1:02d}" if len(months) == 1 else f"{y1}.{m1:02d} ~ {y2}.{m2:02d}"
    return f"{span} · {len(months)}개월 {day_count(h)}일 기록됨"


# ----------------------------------------------------------------------------
# 조회/통계
# ----------------------------------------------------------------------------

def iter_meals(h: dict | None = None):
    """(날짜, 끼니, 6개 항목) 을 날짜순으로 훑는다."""
    h = load() if h is None else h
    for k in sorted(h):
        p = parse_mkey(k)
        if not p or not isinstance(h[k], dict):
            continue
        y, mo = p
        for ds in sorted(h[k], key=lambda s: int(s) if str(s).isdigit() else 0):
            if not str(ds).isdigit():
                continue
            try:
                d = date(y, mo, int(ds))
            except ValueError:
                continue
            secs = h[k][ds]
            if not isinstance(secs, dict):
                continue
            for sec in SECTIONS:
                items = secs.get(sec)
                if items:
                    yield d, sec, _pad(items)


def slot_names(h: dict | None = None) -> dict[str, list[str]]:
    """과거에 실제로 쓴 메뉴 이름을 자리별로 모은다(많이 쓴 순)."""
    h = load() if h is None else h
    cnt = {"mains": Counter(), "sides": Counter(), "soups": Counter()}
    for _d, _sec, items in iter_meals(h):
        if items[IDX_MAIN]:
            cnt["mains"][items[IDX_MAIN]] += 1
        for i in IDX_SIDES:
            if items[i]:
                cnt["sides"][items[i]] += 1
        if items[IDX_SOUP]:
            cnt["soups"][items[IDX_SOUP]] += 1
    return {k: [n for n, _c in c.most_common()] for k, c in cnt.items()}


def all_names(h: dict | None = None) -> set[str]:
    """과거에 쓴 모든 메뉴 이름(김치류 자리는 제외) — ‘과거 메뉴만’ 검증용."""
    s = slot_names(h)
    return {n for k in ("mains", "sides", "soups") for n in s[k] if n}


def main_stats(before: date | None = None, h: dict | None = None) -> dict[str, dict]:
    """메인 요리별 {총 횟수, 마지막 등장일}. before 가 있으면 그 날 이전만 센다."""
    out: dict[str, dict] = {}
    for d, _sec, items in iter_meals(h):
        if before and d >= before:
            continue
        name = items[IDX_MAIN]
        if not name:
            continue
        e = out.setdefault(name, {"count": 0, "last": d})
        e["count"] += 1
        if d > e["last"]:
            e["last"] = d
    return out


def recent_tail(year: int, month: int, days: int = TAIL_DAYS,
                h: dict | None = None) -> list[tuple[int, str, str, str]]:
    """대상 월 직전 며칠의 기록 — (상대일, 끼니, 메인, 날짜표기).

    상대일은 '대상 월 1일 = 1' 기준이라 생성 결과의 일(day) 과 바로 뺄 수 있다.
    (지난달 마지막 날 = 0, 그 전날 = -1 ...)  달이 바뀔 때 같은 메인이 바로
    이어지는 걸 잡아내려고 menu_planner.validate_menu 에 넘긴다.
    """
    first = date(year, month, 1)
    out = []
    for d, sec, items in iter_meals(h):
        rel = (d - first).days + 1
        if rel > 0 or rel <= -days:
            continue
        if items[IDX_MAIN]:
            out.append((rel, sec, items[IDX_MAIN], f"{d.month}/{d.day}"))
    return out


# ----------------------------------------------------------------------------
# 프롬프트 블록
# ----------------------------------------------------------------------------

def _month_text(year: int, month: int, h: dict) -> str:
    return layouts.month_data_to_text(get_month(year, month, h))


def prompt_block(year: int, month: int, mode: str = MODE_REF,
                 months: int = PROMPT_MONTHS) -> str:
    """생성 프롬프트에 실을 과거 기록 블록(모드가 off거나 기록이 없으면 빈 문자열)."""
    if mode == MODE_OFF:
        return ""
    h = load()
    if not h:
        return ""
    first = date(year, month, 1)
    keys = [k for k in sorted(h)
            if (p := parse_mkey(k)) and date(p[0], p[1], 1) < first and h.get(k)]
    if not keys:
        return ""

    parts = [
        "[📚 과거 실제 식단 기록]",
        "아래는 이 가게에서 실제로 냈던 식단입니다. 손님이 익숙한 조합·표기이므로 "
        "같은 결(재료·조리법·구성·이름 표기)로 편성하세요.",
    ]
    for k in keys[-months:]:
        y2, m2 = parse_mkey(k)
        parts.append(f"\n({y2}년 {m2}월 실제로 낸 식단)")
        parts.append(_month_text(y2, m2, h))

    stats = main_stats(before=first, h=h)
    if stats:
        rows = sorted(stats.items(), key=lambda kv: (-kv[1]["count"], kv[0]))
        lines = [f"- {n}: 총 {v['count']}회, 마지막 {v['last'].month}/{v['last'].day}"
                 for n, v in rows[:PROMPT_STAT_ROWS]]
        parts.append("\n[과거 메인 요리 사용 현황]")
        parts.append("\n".join(lines))
        parts.append("→ 마지막에 낸 지 얼마 안 된 메인은 이번 달 앞부분에 넣지 말고, "
                     "총 횟수가 적은 메인부터 채워 골고루 돌아가게 하세요.")

    if mode == MODE_ONLY:
        s = slot_names(h)
        parts.append("\n[⚠ 과거에 냈던 메뉴만 사용]")
        parts.append("새 메뉴를 지어내지 말고, 아래 목록 안에서만 골라 편성하세요. "
                     "목록에 없는 이름은 한 개도 쓰면 안 됩니다.")
        parts.append("· 메인: " + ", ".join(s["mains"]))
        parts.append("· 반찬: " + ", ".join(s["sides"]))
        parts.append("· 국: " + ", ".join(s["soups"]))
    return "\n".join(parts)


# ----------------------------------------------------------------------------
# 가져오기 — 예전에 쓰던 엑셀 식단표에서 한 번에 읽어 오기
# ----------------------------------------------------------------------------

_SHEET_YM = re.compile(r"(\d{2,4})\s*[.\-_/]\s*(\d{1,2})")


def _ym_from_sheet(title: str) -> tuple[int, int] | None:
    """'배달-26.07' / '주방-26.07' 같은 시트 이름에서 연/월을 읽는다."""
    m = _SHEET_YM.search(title or "")
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    if not 1 <= mo <= 12:
        return None
    return (2000 + y if y < 100 else y), mo


def _cell_items(v) -> list[str]:
    """메뉴 칸('메인\n\n반찬\n\n…') → 항목 리스트. 안내문·빈칸은 걸러낸다."""
    if not isinstance(v, str):
        return []
    items = [x.strip() for x in v.split("\n")]
    return [x for x in items if x and not x.startswith("※")]


def import_xlsx(path: str, year: int | None = None,
                month: int | None = None) -> list[tuple[int, int, int]]:
    """엑셀 식단표(.xlsx)를 읽어 기록에 넣는다. [(연, 월, 저장한 날 수), ...] 반환.

    배달형(달력)·주방형(주간) 두 서식을 모두 읽는다. 시트 이름에 연월이 없으면
    인자로 받은 연/월을 쓴다. 기존 기록은 지우지 않고 덮어쓴다(merge).
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    done: list[tuple[int, int, int]] = []
    for ws in wb.worksheets:
        ym = _ym_from_sheet(ws.title) or ((year, month) if year and month else None)
        if not ym:
            continue
        md = MonthData(year=ym[0], month=ym[1])
        section = ""
        days_by_col: dict[int, int] = {}

        for row in ws.iter_rows():
            vals = [c.value for c in row]

            # 제목 칸("7월 식단표 (중식)") — 배달형은 여기서 끼니가 정해진다
            head = next((v.strip() for v in vals if isinstance(v, str) and v.strip()), "")
            if "식단표" in head:
                if "중식" in head:
                    section = "중식"
                elif "석식" in head:
                    section = "석식"
                continue

            # 날짜 행: 1~31 정수가 여러 칸
            nums = {i: int(v) for i, v in enumerate(vals)
                    if isinstance(v, (int, float)) and float(v).is_integer()
                    and 1 <= int(v) <= 31}
            if len(nums) >= 2:
                days_by_col = nums
                continue

            # 구분 칸(주방형): 그 행의 끼니
            if isinstance(vals[0], str) and vals[0].strip() in SECTIONS:
                section = vals[0].strip()

            if not (days_by_col and section):
                continue
            for ci, day in days_by_col.items():
                items = _cell_items(vals[ci] if ci < len(vals) else None)
                if len(items) < 2:      # 안내문·요일 머리글 등은 걸러진다
                    continue
                md.meals.setdefault(day, {})[section] = _pad(items)

        if md.meals:
            done.append((md.year, md.month, save_month(md, merge=True)))
    return done
