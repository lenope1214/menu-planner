# -*- coding: utf-8 -*-
"""
엑셀 서식(exporters.py) 잠금
=============================
메뉴 열은 넓게 고정하고, 행 높이는 들어간 글자 수에 맞춰 계산한다.
**열 너비를 줄이면 메뉴 글자가 잘리고, wrap_text 를 끄면 6개가 한 줄로 붙어 보인다.**
둘 다 눈으로만 확인되는 종류라 여기서 못 박아 둔다.
"""

from __future__ import annotations

from openpyxl import load_workbook

import exporters as ex
from layouts import PRESET_CALENDAR, PRESET_KITCHEN
from helpers import find_cell, line, text

MENU = ["제육볶음", "어묵볶음", "감자조림", "오이무침", "김치류", "콩나물국"]
SAMPLE = text(lunch=[line("7/1(수)", "제육볶음")],
              dinner=[line("7/1(수)", "고등어구이")])


def _sheet(tmp_path, preset):
    p = tmp_path / "t.xlsx"
    ex.save_xlsx(SAMPLE, str(p), preset=preset, year=2026, month=7)
    wb = load_workbook(str(p))
    return wb, wb.worksheets[0]


# ---- 서식 상수 ----

def test_서식_수치를_지킨다():
    assert ex.MENU_COL_WIDTH == 26.0      # 배달 A~G / 주방 B~H 공통
    assert ex.DEFAULT_COL_W == 12.63      # 시트 기본 너비
    assert ex.DLV_ROW_H_HEAD == 37.5      # 배달 제목/요일/날짜 행
    assert ex.KIT_ROW_H_HEAD == 33.75     # 주방 요일/날짜 행


def test_메뉴_여섯개는_열한_줄로_본다():
    """항목 6개 + 사이 빈 줄 5개 = 11줄. 이 너비에서는 접히지 않아야 한다."""
    cell = "\n\n".join(MENU)
    assert ex._cell_lines(cell, ex.MENU_COL_WIDTH, 26) == 11


def test_긴_메뉴명은_접힌_줄까지_센다():
    short = ex._cell_lines("제육볶음", ex.MENU_COL_WIDTH, 26)
    long = ex._cell_lines("매운돼지갈비찜간장양념구이", ex.MENU_COL_WIDTH, 26)
    assert short == 1
    assert long > short


def test_행_높이는_줄_수에_비례하고_한계를_넘지_않는다():
    one = ex.calc_row_height(["제육볶음"], ex.MENU_COL_WIDTH, 26)
    six = ex.calc_row_height(["\n\n".join(MENU)], ex.MENU_COL_WIDTH, 26)
    assert six > one
    assert one >= ex._MIN_ROW_PT                       # 너무 납작해지지 않는다
    huge = ex.calc_row_height(["가" * 2000], ex.MENU_COL_WIDTH, 26)
    assert huge == ex._MAX_ROW_PT                      # 과도한 줄바꿈은 잘라낸다


# ---- 배달형(달력) ----

def test_배달형_시트와_열너비(tmp_path):
    wb, ws = _sheet(tmp_path, PRESET_CALENDAR)
    assert wb.sheetnames == ["배달-26.07"]
    for c in "ABCDEFG":
        assert ws.column_dimensions[c].width == ex.MENU_COL_WIDTH
    assert ws.page_setup.orientation == "portrait"


def test_넓힌_열은_인쇄시_한_페이지_폭에_맞춘다(tmp_path):
    """열을 넓혔으므로 fitToWidth 가 없으면 인쇄가 옆으로 잘린다."""
    _wb, ws = _sheet(tmp_path, PRESET_CALENDAR)
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0


def test_배달형_메뉴칸은_줄바꿈이_켜져_있다(tmp_path):
    """원본은 구글 시트라 wrapText 없이도 보이지만, 엑셀은 꺼져 있으면 한 줄로 붙는다."""
    _wb, ws = _sheet(tmp_path, PRESET_CALENDAR)
    cell = find_cell(ws, "제육볶음")
    assert cell is not None
    assert cell.alignment.wrap_text is True
    assert cell.value == "\n\n".join(MENU)      # 항목 사이 빈 줄
    assert ws.row_dimensions[cell.row].height == ex.calc_row_height(
        [cell.value], ex.MENU_COL_WIDTH, 26)


def test_배달형_요일머리글은_전체표기(tmp_path):
    _wb, ws = _sheet(tmp_path, PRESET_CALENDAR)
    assert find_cell(ws, "일요일") is not None
    assert find_cell(ws, "토요일") is not None


def test_안내문이_들어간다(tmp_path):
    _wb, ws = _sheet(tmp_path, PRESET_CALENDAR)
    assert find_cell(ws, "※") is not None


# ---- 주방형(주간) ----

def test_주방형_시트와_열너비(tmp_path):
    wb, ws = _sheet(tmp_path, PRESET_KITCHEN)
    assert wb.sheetnames == ["주방-26.07"]
    assert ws.column_dimensions["A"].width == 13.0      # '구분' 라벨 열
    for c in "BCDEFGH":
        assert ws.column_dimensions[c].width == ex.MENU_COL_WIDTH
    assert ws.page_setup.orientation == "landscape"


def test_주방형_메뉴칸과_구분칸(tmp_path):
    _wb, ws = _sheet(tmp_path, PRESET_KITCHEN)
    cell = find_cell(ws, "제육볶음")
    assert cell.alignment.wrap_text is True
    assert ws.row_dimensions[cell.row].height == ex.calc_row_height(
        [cell.value], ex.MENU_COL_WIDTH, 24)
    assert ws.cell(row=cell.row, column=1).value == "중식"


# ---- 메뉴명이 길어지면 행만 키운다 ----

def test_긴_메뉴명이면_행_높이가_커진다(tmp_path):
    long_text = text(lunch=[line("7/1(수)", "매운돼지갈비찜간장양념구이")])
    p = tmp_path / "long.xlsx"
    ex.save_xlsx(long_text, str(p), preset=PRESET_CALENDAR, year=2026, month=7)
    ws = load_workbook(str(p)).worksheets[0]
    cell = find_cell(ws, "매운돼지갈비찜간장양념구이")

    _wb2, ws_short = _sheet(tmp_path, PRESET_CALENDAR)
    short_cell = find_cell(ws_short, "제육볶음")
    assert (ws.row_dimensions[cell.row].height
            > ws_short.row_dimensions[short_cell.row].height)


# ---- 텍스트 저장 ----

def test_텍스트로도_저장된다(tmp_path):
    p = tmp_path / "t.txt"
    ex.save_txt(SAMPLE, str(p))
    assert p.read_text(encoding="utf-8") == SAMPLE
